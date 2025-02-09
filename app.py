from flask import Flask, render_template, request, redirect, url_for, send_file, session
from docx import Document
from utils import extract_placeholders, fill_data
import os

app = Flask(__name__)
app.secret_key = 'your_secret_key'  # Needed for session handling
app.config['UPLOAD_FOLDER'] = 'uploads'  # Folder to store uploaded files
app.config['DOC_FOLDER'] = 'docs'        # Folder to store generated .docx files

# Ensure upload and doc folders exist
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
os.makedirs(app.config['DOC_FOLDER'], exist_ok=True)


@app.route('/', methods=['GET', 'POST'])
def setup():
    if request.method == 'POST':
        # Handle the uploaded template file
        uploaded_file = request.files.get('template_filename')
        if uploaded_file and uploaded_file.filename:
            template_filename = uploaded_file.filename
            template_path = os.path.join(app.config['UPLOAD_FOLDER'], template_filename)
            uploaded_file.save(template_path)

            # Store the file name in the session
            session['template_filename'] = template_filename
        else:
            return "Please upload a valid .docx template."

        # Get the output document filename
        session['doc_filename'] = request.form.get('doc_filename')

        # Redirect to the form page
        return redirect(url_for('form'))

    return render_template('setup.html')


@app.route('/form', methods=['GET', 'POST'])
def form():
    template_filename = session.get('template_filename')
    doc_filename = session.get('doc_filename')

    # Check if filenames were provided
    if not template_filename or not doc_filename:
        return redirect(url_for('setup'))
    
    # Load the template file
    template_path = os.path.join(app.config['UPLOAD_FOLDER'], template_filename)
    try:
        doc = Document(template_path)
    except Exception as e:
        return f"Error loading template: {e}"

    # Extract placeholders from the first column of each table
    placeholders = extract_placeholders(doc)

    if request.method == 'POST':
        # Collect user input
        form_data = {key: request.form.get(key) for key in placeholders}

        # Handle uploaded files
        uploaded_files = request.files.getlist('uploaded_files')
        for file in uploaded_files:
            if file.filename:
                file_path = os.path.join(app.config['UPLOAD_FOLDER'], file.filename)
                file.save(file_path)

        # Fill the data into the document
        fill_data(doc, form_data)

        # Save the filled document
        doc_path = os.path.join(app.config['DOC_FOLDER'], doc_filename)
        doc.save(doc_path)

        return send_file(doc_path, as_attachment=True)
    
    # Pass the placeholders to the HTML template
    return render_template('index.html', placeholders=placeholders)

if __name__ == "__main__":
    app.run(debug=True)
